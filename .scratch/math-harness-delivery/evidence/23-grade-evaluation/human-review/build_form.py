"""生成人工判定表：严重度、终稿八维评分、发现是否成立。

表中不出现样本答案、评阅者身份和模型给出的严重度；这些写入被 Git 忽略的
work/grade-evaluation/human-review/mapping.json，判定结果回收后再合并入库。
用法：uv run --with openpyxl python .scratch/math-harness-delivery/evidence/23-grade-evaluation/human-review/build_form.py
"""

import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from teaching_harness.grade_evaluation import cli
from teaching_harness.grade_evaluation.calibration import SampleAnswer, score_detection
from teaching_harness.grade_evaluation.checks import TEACHER_QUESTIONS
from teaching_harness.grade_evaluation.evidence import pointer_parts
from teaching_harness.grade_evaluation.runs import CalibrationReport

HERE = Path(__file__).parent
OUT = HERE / "judgment-form.xlsx"
MAPPING = cli.WORK / "human-review/mapping.json"
SEED = 23
REVIEWERS = ("r1", "r2", "rd")
QUOTAS = {"rd": 15, "r1": 5, "r2": 5}
SEVERITY_ZH = {"critical": "重大失败", "key_gap": "关键缺口", "local": "局部问题"}
FIELD = {
    "narrative": "叙述",
    "entry": "进入条件",
    "exit": "完成表现",
    "assessment_plan": "评价安排",
    "title": "标题",
    "lesson_count": "课时",
    "prerequisite_units": "先备单元",
    "opportunity": "学习机会",
    "evidence": "达成证据",
    "unit_id": "单元",
    "handoff_guidance": "交接说明",
    "design_inferences": "设计推断",
    "reserve_plan": "机动计划",
    "reserve_lessons": "机动课时",
    "teacher_preparation": "教师准备",
    "prerequisites": "先备说明",
    "successors": "后续说明",
    "student_actions": "学生行为",
    "unit_ids": "承担单元",
    "solution": "作者解答",
    "design_consequence": "设计结论",
    "support": "支持语",
    "prompt": "题面",
    "blocks": "题面块",
    "anticipated_response": "预判回应",
    "focus_unit_id": "焦点单元",
    "assumptions": "假设",
    "limitations": "局限",
    "knowledge_uses": "知识依据",
}

# A：严重度。前 10 项是第二版保留集的注入问题，后 3 项是终稿上被检查确认或反复报告的问题。
A_ITEMS = [
    {
        "sample": "h2-01",
        "answer": "h2i-01",
        "check": "8.F.B.5 在全年中的学习机会与达成证据，是否覆盖该标准的全部要求",
        "excerpt": (
            "【标准 8.F.B.5】Describe qualitatively the functional relationship between two "
            "quantities by analyzing a graph (e.g., where the function is increasing or "
            "decreasing, linear or nonlinear). Sketch a graph that exhibits the qualitative "
            "features of a function that has been described verbally.\n"
            "【候选中 8.F.B.5 的全部分配】仅 1 条，单元 3（函数）讲授：\n"
            "学习机会：在函数图象上选取两个不同点计算平均变化率，依据差商的正负判断函数在指定区间内的增减，"
            "并用计算结果比较两段折线的倾斜程度。\n"
            "达成证据：检查学生能否从给定折线图中读出两点坐标、准确计算差商，并依据差商的正负与大小写出增减结论。\n"
            "（全文其他位置没有 8.F.B.5 的安排。）"
        ),
    },
    {
        "sample": "h2-03",
        "answer": "h2i-02",
        "check": "全年规划对八项数学实践（MP1–MP8）的安排",
        "excerpt": (
            "【候选列出的数学实践】MP1、MP2、MP3、MP4、MP5、MP6、MP8，各有承担单元、学生行为与观察证据。\n"
            "（候选中没有 MP7 的条目；请求要求说明目标与数学实践的实际学习机会。）"
        ),
    },
    {
        "sample": "h2-04",
        "answer": "h2i-03",
        "check": "单元 4（二元一次方程组）叙述中，两条直线的位置关系与方程组解的对应",
        "excerpt": (
            "【单元 4 叙述摘录】……单元首先建立解即为两直线几何交点的直观，随后系统学习代入消元法与加减消元法，"
            "并借助直线重合的几何形态理解方程组无解、借助直线平行的几何形态理解方程组无穷多解的代数本质，"
            "最后在方案抉择与成本平衡情境中应用建模。"
        ),
    },
    {
        "sample": "h2-05",
        "answer": "h2i-04",
        "check": "单元的教学顺序与各单元声明的先备关系",
        "excerpt": (
            "【教学顺序（按排列）】1. 线性方程与斜率（unit_2）→ 2. 几何变换与相似（unit_1）→ 3. 函数 → "
            "4. 方程组 → 5. 实数与勾股 → 6. 指数 → 7. 体积 → 8. 双变量统计\n"
            "【unit_2 的先备单元】unit_1_geom_transform\n"
            "【unit_2 叙述摘录】本单元承接 Unit 1 的相似三角形理论，在平面直角坐标系中构造‘斜率三角形’，"
            "严密证明非垂直直线上任意两不同点的变化率恒定……"
        ),
    },
    {
        "sample": "h2-06",
        "answer": "h2i-05",
        "check": "8.SP.A.4（双向列联表）的评价安排与学习机会的先后",
        "excerpt": (
            "【教学顺序】单元 1–8 按原序；方程组为单元 4（第一学期末），双变量统计为单元 8（第二学期末）。\n"
            "【8.SP.A.4 的分配】\n"
            "① 单元 4 评价：在第一学期期末的方程组单元测试中加入双向列联表题组，要求学生由频数表计算行、列相对频率"
            "并判断两个分类变量是否相关。\n"
            "② 单元 8 讲授：探究双变量分类数据，编制并解读双向列联表……"
        ),
    },
    {
        "sample": "h2-07",
        "answer": "h2i-06",
        "check": "各单元课时、机动课时与学年总课时",
        "excerpt": (
            "【学校条件】全年 180 节，其中机动 20 节。\n"
            "【各单元课时】22、22、22、18、26、36（指数单元）、14、20，合计 180。\n"
            "【机动计划】全年设置 20 节机动课时……不计入各单元常规教学课时。\n"
            "【全年叙述】全课程安排 160 节常规课时与 20 节依据诊断干预的机动课时。"
        ),
    },
    {
        "sample": "h2-09",
        "answer": "h2i-07",
        "check": "单元 7（体积）的评价方式与学校条件",
        "excerpt": (
            "【单元 7 评价安排】学生课后在家使用量杯和水完成等底等高圆锥与圆柱的注水对比实验，并用手机录制实验视频提交，"
            "作为评估体积公式几何关系的主要依据；课内随堂练习（利用普通计算器）评估多步公式代入与根号近似计算；"
            "单元末闭卷测试……所有必需评价均在课内完成。\n"
            "【学校条件】资源：纸笔、直尺、方格纸、黑白打印、教师投影、普通计算器。不假定学生有联网设备；"
            "必要学习与评价在课内，课后练习可选。"
        ),
    },
    {
        "sample": "h2-10",
        "answer": "h2i-08",
        "check": "8.F.B.4 在全年叙述中的后续应用与目标分配",
        "excerpt": (
            "【8.F.B.4 的分配】仅 1 条：单元 3（函数）讲授。\n"
            "【全年叙述】最后在 Unit 8，学生运用全年所学的线性模型（8.F.B.4 进阶应用）对双变量合成数据进行经验拟合与预测分析……\n"
            "【设计推断】8.F.B.4 作为全年初中函数建模的核心……后启 Unit 8 的双变量统计散点图经验拟合……"
        ),
    },
    {
        "sample": "h2-12",
        "answer": "h2i-09",
        "check": "坐标距离探查的设计结论所依据的事实",
        "excerpt": (
            "【探查 task_3 设计结论摘录】……在本校两个八年级班级共 56 名学生的课内试教中，49 人独立完成了直角顶点定位"
            "与距离计算，表明学生已掌握该方法，Unit 5 前无需再安排相应的进入诊断。\n"
            "【学校条件】无班级诊断、个人掌握度或已确认支持记录。"
        ),
    },
    {
        "sample": "h2-14",
        "answer": "h2i-10",
        "check": "焦点单元与交接对象是否一致",
        "excerpt": (
            "【请求】提供围绕 8.F.B.4 的实际目标单元交接，但不预设单元编号。\n"
            "【focus_unit_id】unit_4_linear_systems（二元一次方程组）\n"
            "【交接说明开头】【8.F.B.4 核心目标单元（Unit 3: 函数概念与线性函数建模）交接指引】"
            "……8.F.B.4 在本单元承担最为核心的明确讲授（teach）职责……\n"
            "【8.F.B.4 的分配】单元 3 讲授、单元 8 应用；方程组单元没有 8.F.B.4 的分配。"
        ),
    },
    {
        "sample": "b-final",
        "answer": None,
        "check": "交接说明中 8.F.B.4 在单元 4 的应用与目标分配",
        "excerpt": (
            "【交接说明摘录】本单元建构的 8.F.B.4 建模技能将在 Unit 4（二元一次方程组的交点平衡状态）中被应用，"
            "并在 Unit 8（双变量统计）中作为散点图经验拟合直线的分析工具（8.SP.A.3）实施迁移应用。\n"
            "【8.F.B.4 的分配】单元 3 讲授、单元 8 应用；单元 4 没有 8.F.B.4 的分配。"
        ),
    },
    {
        "sample": "b-final",
        "answer": None,
        "check": "8.F.A.1 学习机会中对函数图象的表述",
        "excerpt": (
            "【标准 8.F.A.1】Understand that a function is a rule that assigns to each input exactly "
            "one output. The graph of a function is the set of ordered pairs consisting of an "
            "input and the corresponding output.\n"
            "【候选学习机会】通过自动售票计价、输入输出数值对应等实例，建立函数的统一定义：每个输入值对应唯一确定输出值的规则；"
            "理解图象是所有有序实数对 $(x, y)$ 的集合。"
        ),
    },
    {
        "sample": "b-final",
        "answer": None,
        "check": "单元 5 叙述中勾股定理及其逆定理的证明方式",
        "excerpt": (
            "【单元 5 叙述摘录】……随后通过面积割补法证明勾股定理及其逆定理，熟练求解 x^2=p 与 x^3=p，"
            "并将勾股定理推广应用于二维平面……"
        ),
    },
]

A_HEAD = ["编号", "样本", "核对点", "原文与事实摘录", "是否存在问题", "严重度", "判断依据", "把握"]
B_HEAD = [
    "编号",
    "维度",
    "权重",
    "审查范围",
    "锚点（0 / 2 / 4）",
    "分数",
    "是否有重大失败",
    "判断依据",
    "把握",
]
C_HEAD = [
    "编号",
    "样本",
    "位置",
    "原文引文",
    "发现陈述",
    "是否成立",
    "若成立，严重度",
    "判断依据",
    "把握",
]

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=13, bold=True)
EXAMPLE = Font(name="Arial", size=10, italic=True, color="808080")
INPUT = PatternFill("solid", fgColor="FFF2CC")
HEADER = PatternFill("solid", fgColor="D9E2F3")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def sample_label(sample_id: str) -> str:
    return "终稿" if sample_id == "b-final" else sample_id.upper()


def locate(locator: str, document_id: str, content: dict) -> str:
    if document_id.endswith(".standards"):
        return "年级标准原文"
    if document_id.endswith(".conditions"):
        return "学校条件"
    parts = pointer_parts(locator)
    if not parts:
        return "全文"
    head = parts[0]
    try:
        if head == "units":
            unit = content["units"][int(parts[1])]
            field = FIELD.get(parts[2], parts[2]) if len(parts) > 2 else ""
            return f"单元 {unit['id']} · {field}".rstrip(" ·")
        if head == "goals":
            goal = content["goals"][int(parts[1])]
            if len(parts) > 3:
                a = goal["allocations"][int(parts[3])]
                field = FIELD.get(parts[4], parts[4]) if len(parts) > 4 else ""
                return f"目标 {goal['code']} · {a['unit_id']}:{a['role']} · {field}".rstrip(" ·")
            return f"目标 {goal['code']}"
        if head == "tasks":
            task = content["tasks"][int(parts[1])]
            field = FIELD.get(parts[2], parts[2]) if len(parts) > 2 else ""
            return f"探查 {task['id']} · {field}".rstrip(" ·")
        if head == "practices":
            practice = content["practices"][int(parts[1])]
            field = FIELD.get(parts[2], parts[2]) if len(parts) > 2 else ""
            return f"数学实践 {practice['code']} · {field}".rstrip(" ·")
    except (IndexError, KeyError, ValueError):
        pass
    return FIELD.get(head, head) + ("" if len(parts) == 1 else f"（{locator}）")


def expected_hits(sample, answer, issue, sources):
    """某个预期问题在各来源中被命中时的最高严重度。"""
    rank = {"local": 0, "key_gap": 1, "critical": 2}
    one = SampleAnswer(
        sample_id=sample.id,
        kind=answer.kind,
        expected=[issue],
        protected_pointers=[],
        rationale=answer.rationale,
    )
    result = {}
    for name, (findings, evidence) in sources.items():
        records = {e.id: e for e in evidence}
        best = None
        for f in findings:
            [m] = score_detection(
                [sample],
                [one],
                {sample.id: [f]},
                records,
                "model",
                origins={"model", "program"},
                group="x",
            )
            if m.detected and (best is None or rank[f.severity] > rank[best]):
                best = f.severity
        result[name] = best
    return result


def main() -> None:
    answers = {a.sample_id: a for a in cli.answers()}
    rubric = json.loads(cli.RUBRIC.read_text())
    program = cli.PROGRAM.validate_json((cli.RESULTS / "program.json").read_bytes())
    stages = {s: cli._stage_runs(s) for s in cli.STAGES}
    runs = {r: cli._runs(r) for r in REVIEWERS}
    calibration = CalibrationReport.model_validate_json(
        (cli.RESULTS / "calibration.json").read_bytes()
    )
    per_sample = calibration.sets["holdout-v2"].per_sample
    mapping: dict = {"seed": SEED, "A": [], "B": [], "C": []}

    wb = Workbook()
    guide = wb.active
    guide.title = "说明"

    # A
    a_sheet = wb.create_sheet("A严重度")
    a_rows = []
    for n, item in enumerate(A_ITEMS, start=1):
        code = f"A{n:02d}"
        a_rows.append([code, sample_label(item["sample"]), item["check"], item["excerpt"]])
        record = {"item": code, "sample": item["sample"], "answer": item["answer"]}
        if item["answer"]:
            sample = cli.sample_by_id(item["sample"])
            answer = answers[sample.id]
            issue = next(e for e in answer.expected if e.id == item["answer"])
            sources = {"program": (program[sample.id].findings, program[sample.id].evidence)}
            for stage, stage_runs in stages.items():
                sources[stage] = (
                    stage_runs[sample.id].findings(),
                    stage_runs[sample.id].evidence(),
                )
            for r in REVIEWERS:
                sources[r] = (runs[r][sample.id].findings, runs[r][sample.id].evidence)
            record["answer_severity"] = issue.severity
            record["detected"] = expected_hits(sample, answer, issue, sources)
        mapping["A"].append(record)
    fill_sheet(
        a_sheet,
        "A　严重度判定（13 项）",
        "逐项阅读“原文与事实摘录”，判断“核对点”处是否存在问题；若存在，按“说明”页的定义给出严重度。",
        A_HEAD,
        [
            "示例",
            "H2-00",
            "单元 X 的某段表述",
            "（摘录）",
            "存在",
            "关键缺口",
            "摘录中……与标准要求……不符，会改变……",
            "中",
        ],
        a_rows,
        widths=[7, 9, 26, 70, 12, 12, 40, 8],
        inputs=[4, 5, 6, 7],
        lists={
            4: "存在,部分存在,不存在",
            5: "重大失败,关键缺口,局部问题,不构成问题",
            7: "高,中,低",
        },
    )

    # B
    b_sheet = wb.create_sheet("B终稿评分")
    b_rows = []
    for c in rubric["criteria"]:
        anchors = "\n".join(f"{k} 分：{v}" for k, v in c["anchors"].items())
        scope = c["review_scope"]
        if c["id"] == "Q8":
            scope += "\n教师规划问题：" + "".join(TEACHER_QUESTIONS.values())
        b_rows.append([c["id"], c["name"], c["weight"], scope, anchors])
        scores = {}
        for r in REVIEWERS:
            rating = next(
                (x for x in runs[r]["b-final"].ratings if x.criterion_id == c["id"]), None
            )
            scores[r] = rating.score if rating else None
        mapping["B"].append({"item": c["id"], "model_scores": scores})
    fill_sheet(
        b_sheet,
        "B　15 终稿八维评分（8 项）",
        "请先通读终稿阅读版（curriculum.html），再独立打分；本表不提供任何模型的分数或评语。",
        B_HEAD,
        [
            "示例",
            "某维度",
            10,
            "（审查范围）",
            "（锚点）",
            "3",
            "否",
            "应查范围已核完，主要联系有具体证据；距 4 分的差距是……",
            "中",
        ],
        b_rows,
        widths=[7, 16, 6, 34, 60, 8, 12, 44, 8],
        inputs=[5, 6, 7, 8],
        lists={5: "0,1,2,3,4,未定", 6: "是,否", 8: "高,中,低"},
    )

    # C
    rng = random.Random(SEED)
    avoid = ("逆定理", "有序实数对", "Unit 4", "unit_4")
    picked = []
    for r in REVIEWERS:
        pool = []
        scope = [s.id for s in cli.sample_sets()["holdout-v2"]] + ["b-final"]
        for sid in sorted(set(runs[r]) & set(scope)):
            run = runs[r][sid]
            if sid == "b-final":
                allowed = {f.id for f in run.findings}
            else:
                m = per_sample[r][sid]
                allowed = set(m.other_findings) | set(m.false_positives)
            for f in run.findings:
                text = f"{f.claim}{f.requirement}"
                if f.id in allowed and not any(k in text for k in avoid):
                    pool.append((r, sid, f, run))
        picked += rng.sample(pool, QUOTAS[r])
    rng.shuffle(picked)
    c_sheet = wb.create_sheet("C发现核对")
    c_rows = []
    for n, (r, sid, f, run) in enumerate(picked, start=1):
        code = f"C{n:02d}"
        content = json.loads((cli.EVAL / f"candidates/{sid}.json").read_text())
        records = {e.id: e for e in run.evidence}
        cited = [records[e] for e in f.evidence_ids if e in records][:2]
        where = "；".join(dict.fromkeys(locate(e.locator, e.document_id, content) for e in cited))
        quotes = "\n".join(f"「{e.quote[:260]}」" for e in cited)
        statement = f"{f.claim}\n（依据的要求：{f.requirement}）"
        c_rows.append([code, sample_label(sid), where, quotes, statement])
        mapping["C"].append(
            {
                "item": code,
                "reviewer": r,
                "sample": sid,
                "finding": f.id,
                "model_severity": f.severity,
            }
        )
    fill_sheet(
        c_sheet,
        "C　发现是否成立（25 项）",
        "每项是某次模型评阅报告的一条发现（不标明来自哪个模型）。请对照引文和终稿阅读版，判断它是否成立；若成立，给出严重度。"
        "H2 开头的样本与终稿只差一处改动，改动内容见 A 页对应样本。",
        C_HEAD,
        [
            "示例",
            "终稿",
            "单元 X · 叙述",
            "「引文」",
            "某项安排缺少……",
            "部分成立",
            "局部问题",
            "引文属实，但……不影响年度决定",
            "高",
        ],
        c_rows,
        widths=[7, 9, 24, 50, 50, 12, 12, 40, 8],
        inputs=[5, 6, 7, 8],
        lists={5: "成立,部分成立,不成立", 6: "重大失败,关键缺口,局部问题", 8: "高,中,低"},
    )

    write_guide(guide, rubric, len(a_rows), len(b_rows), len(c_rows))
    # 只有完成进度用到公式；打开时由表格软件重算。
    wb.calculation.fullCalcOnLoad = True
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    MAPPING.parent.mkdir(parents=True, exist_ok=True)
    MAPPING.write_text(json.dumps(mapping, ensure_ascii=False, indent=2) + "\n")
    print(OUT, MAPPING, len(a_rows), len(b_rows), len(c_rows))


def fill_sheet(ws, title, note, head, example, rows, *, widths, inputs, lists):
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = note
    ws["A2"].font = FONT
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(head))
    ws.row_dimensions[2].height = 30
    for col, value in enumerate(head, start=1):
        cell = ws.cell(row=3, column=col, value=value)
        cell.font, cell.fill, cell.border, cell.alignment = BOLD, HEADER, BOX, WRAP
    for col, value in enumerate(example, start=1):
        cell = ws.cell(row=4, column=col, value=value)
        cell.font, cell.border, cell.alignment = EXAMPLE, BOX, WRAP
    for r, row in enumerate(rows, start=5):
        longest = 0
        for col in range(1, len(head) + 1):
            value = row[col - 1] if col <= len(row) else None
            cell = ws.cell(row=r, column=col, value=value)
            cell.font, cell.border, cell.alignment = FONT, BOX, WRAP
            if col - 1 in inputs:
                cell.fill = INPUT
            if isinstance(value, str):
                lines = sum(
                    len(part) // max(1, int(widths[col - 1] / 1.1)) + 1
                    for part in value.split("\n")
                )
                longest = max(longest, lines)
        ws.row_dimensions[r].height = min(400, max(45, longest * 14))
    last = 4 + len(rows)
    for index, options in lists.items():
        letter = ws.cell(row=5, column=index + 1).column_letter
        rule = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        ws.add_data_validation(rule)
        rule.add(f"{letter}5:{letter}{last}")
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=3, column=col).column_letter].width = width
    ws.freeze_panes = "C4"


def write_guide(ws, rubric, a_count, b_count, c_count):
    lines = [
        ("年级课程检查器：人工判定表", TITLE),
        (
            "用途：校准模型评阅与专项检查的严重度判断、维度评分和发现准确性。你的判断是基准，不需要参考任何模型意见。",
            FONT,
        ),
        ("", FONT),
        ("怎样填写", BOLD),
        (
            "1. 只填浅黄色格子；可从下拉列表选择，“判断依据”写一两句理由，引用原文位置更好。第 4 行是示例，不计入。",
            FONT,
        ),
        (
            "2. 顺序建议：先通读终稿阅读版 curriculum.html（约 40 分钟），做 B 页；再做 A 页和 C 页。合计约 1.5–2 小时。",
            FONT,
        ),
        (
            "3. 拿不准时仍请给出最接近的选项，并把“把握”选为“低”；确实无法判断时，在“判断依据”写明缺什么。",
            FONT,
        ),
        (
            "4. 表中不标明哪条意见来自哪个模型，也不给出样本设计者预设的答案，请按原文独立判断。",
            FONT,
        ),
        ("", FONT),
        ("严重度定义（依据评估协议）", BOLD),
        (
            "重大失败：经原文确认属于以下任一类——"
            + "；".join(rubric["critical_failures"])
            + "。合理的顺序差异、尚未生成逐课材料、语言风格差异不算。",
            FONT,
        ),
        (
            "关键缺口：缺少会改变目标责任、必要依赖、学习机会、评价用途或核心可行性的联系，但不至于使全年进程整体不可执行。",
            FONT,
        ),
        ("局部问题：有安排，但存在不改变年度决定的局部不足或表述不精确。", FONT),
        ("不构成问题 / 不成立：原文没有该问题，或所指情况合理。", FONT),
        ("", FONT),
        ("评分锚点（B 页）", BOLD),
        (
            "0–4 分；各维 0、2、4 分锚点见 B 页。1 分："
            + rubric["intermediate_anchors"]["1"]
            + "。3 分："
            + rubric["intermediate_anchors"]["3"]
            + "。",
            FONT,
        ),
        (
            "维度收敛："
            + "；".join(f"{v}" for v in rubric["dimension_aggregation"].values())
            + "。确实无法判断时选“未定”。",
            FONT,
        ),
        ("", FONT),
        ("随表材料", BOLD),
        (
            "curriculum.html：15 终稿的教师阅读版（与被评 JSON 同源）。H2 样本是终稿副本上的一处改动，改动处已在 A 页摘录。",
            FONT,
        ),
        ("", FONT),
        ("完成进度（自动计算）", BOLD),
    ]
    for r, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if len(text) > 60:
            ws.row_dimensions[r].height = 15 * (len(text) // 60 + 1)
    base = len(lines) + 1
    progress = [
        ("A 严重度", f"=COUNTA('A严重度'!F5:F{4 + a_count})", a_count),
        ("B 终稿评分", f"=COUNTA('B终稿评分'!F5:F{4 + b_count})", b_count),
        ("C 发现核对", f"=COUNTA('C发现核对'!F5:F{4 + c_count})", c_count),
    ]
    ws.cell(row=base, column=1, value="部分").font = BOLD
    ws.cell(row=base, column=2, value="已填").font = BOLD
    ws.cell(row=base, column=3, value="总数").font = BOLD
    for i, (name, formula, total) in enumerate(progress, start=1):
        ws.cell(row=base + i, column=1, value=name).font = FONT
        ws.cell(row=base + i, column=2, value=formula).font = FONT
        ws.cell(row=base + i, column=3, value=total).font = FONT
    ws.column_dimensions["A"].width = 110
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 8


if __name__ == "__main__":
    main()
